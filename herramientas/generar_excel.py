"""Excel comunidad x especialidad a partir de un CSV largo.
Uso: python generar_excel.py plantilla_excel.xlsx plazas_especialidad_2027.csv estado_ccaa_2027.csv salida.xlsx [notas_excel.txt]
datos.csv: fila;comunidad;anio;estado;cuerpo;codigo;especialidad;idioma;plazas;libre;reserva;acceso;aprobadas;extra;nota;fuente
estados.csv: fila;anio;estado;fuente   (una línea por fila de la matriz, aunque no tenga reparto)
"""
import sys, csv, copy, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC, DATA, ESTADOS, OUT = sys.argv[1:5]
import os; os.makedirs(os.path.dirname(os.path.abspath(OUT)), exist_ok=True)
src = openpyxl.load_workbook(SRC)
data = list(csv.DictReader(open(DATA, encoding="utf-8-sig"), delimiter=";"))
estados = list(csv.DictReader(open(ESTADOS, encoding="utf-8-sig"), delimiter=";"))
FILAS = [e["fila"] for e in estados]

thin = Side(style="thin", color="FFBFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def cp(dst, s):
    dst.font = copy.copy(s.font); dst.fill = copy.copy(s.fill); dst.border = copy.copy(s.border)
    dst.alignment = copy.copy(s.alignment); dst.number_format = s.number_format

def key_of(h):
    m = re.match(r"(\d{4})-(\d{3})", h)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    m = re.match(r"(\d{4})-(.*)", h)
    return f"{m.group(1)}-{m.group(2).strip().lower()}"

def row_key(r):
    cu = r["cuerpo"].zfill(4)
    if r["codigo"].strip():
        return f"{cu}-{r['codigo'].strip().zfill(3)}"
    return f"{cu}-{r['especialidad'].strip().lower()}"

ESTADO_FILL = {"Convocada": "FFC6EFCE", "Acordada": "FFDDEBF7", "Acordada en mesa": "FFDDEBF7",
               "OEP publicada": "FFFFF2CC", "Borrador": "FFFCE4D6", "Anuncio": "FFEDEDED"}

out = openpyxl.Workbook(); out.remove(out.active)
pendientes = []

for sname, cuerpos in [("Secundaria", None), ("Maestros", {"0597"})]:
    ws_src = src[sname]
    heads = [(ws_src.cell(1, c).value, ws_src.cell(1, c)) for c in range(6, ws_src.max_column + 1) if ws_src.cell(1, c).value]
    known = {key_of(h) for h, _ in heads}
    rows = [r for r in data if r["plazas"].strip() and r["especialidad"].strip() and
            ((cuerpos and r["cuerpo"].zfill(4) in cuerpos) or (not cuerpos and r["cuerpo"].zfill(4) != "0597"))]
    # columnas nuevas: van tras la última de su cuerpo con código menor
    for r in rows:
        k = row_key(r)
        if k in known:
            continue
        cu = k[:4]
        label = f"{k}-{r['especialidad'].strip()}" if r["codigo"].strip() else f"{cu}-{r['especialidad'].strip()}"
        same = [i for i, (h, _) in enumerate(heads) if h.startswith(cu)]
        if same:
            pos = same[-1]
            if r["codigo"].strip():
                for i in same:
                    if key_of(heads[i][0]) < k:
                        pos = i
            style_cell = heads[same[0]][1]
        else:
            prev = [i for i, (h, _) in enumerate(heads) if h[:4] < cu]
            pos = prev[-1] if prev else len(heads) - 1
            style_cell = heads[pos][1]
        heads.insert(pos + 1, (label, style_cell))
        known.add(k)
    col_of = {key_of(h): 6 + j for j, (h, _) in enumerate(heads)}

    ws = out.create_sheet(sname)
    for i, h in enumerate(["Comunidad", "Año", "Estado", "Fuente", "Total"], 1):
        cp(ws.cell(1, i, h), ws_src.cell(1, 1))
    for j, (h, sc) in enumerate(heads, 6):
        cp(ws.cell(1, j, h), sc)
    last = get_column_letter(5 + len(heads))
    body = ws_src.cell(2, 1)
    plain = ws_src.cell(7, 100) if sname == "Secundaria" else ws_src.cell(2, 6)
    for i, e in enumerate(estados, 2):
        for c in range(1, 6 + len(heads)):
            cp(ws.cell(i, c), body if c <= 5 else plain)
        ws.cell(i, 1, e["fila"])
        mine = [r for r in rows if r["fila"] == e["fila"]]
        vals = {}
        for r in mine:
            vals[row_key(r)] = vals.get(row_key(r), 0) + int(r["plazas"])
        if mine:
            ws.cell(i, 2, int(mine[0]["anio"]))
            ws.cell(i, 3, mine[0]["estado"])
            ws.cell(i, 4, mine[0]["fuente"])
        elif e.get("estado_" + sname, e["estado"]):
            est0 = e.get("estado_" + sname, e["estado"])
            ws.cell(i, 2, int(e["anio"]) if e["anio"] else 2027)
            ws.cell(i, 3, est0 + (" (sin reparto)" if not est0.startswith(("Sin datos", "No convoca")) else ""))
            ws.cell(i, 4, e.get("fuente_" + sname, e["fuente"]))
        est = ws.cell(i, 3).value or ""
        for k2, f in ESTADO_FILL.items():
            if est.startswith(k2):
                ws.cell(i, 3).fill = PatternFill("solid", fgColor=f)
        for k, v in vals.items():
            ws.cell(i, col_of[k], v)
        ws.cell(i, 5, f"=SUM(F{i}:{last}{i})")
        ws.cell(i, 5).font = Font(name="Calibri", size=9, bold=True)
    tr = len(estados) + 2
    tot = ws_src.cell(56, 100) if sname == "Secundaria" else ws_src.cell(126, 6)
    for c in range(1, 6 + len(heads)):
        cp(ws.cell(tr, c), tot)
        if c >= 5:
            L = get_column_letter(c)
            ws.cell(tr, c, f"=SUM({L}2:{L}{tr - 1})")
    ws.cell(tr, 1, "TOTAL")
    for L, w in zip("ABCDE", [20, 6, 24, 34, 7]):
        ws.column_dimensions[L].width = w
    for j in range(len(heads)):
        ws.column_dimensions[get_column_letter(6 + j)].width = 4.43
    ws.row_dimensions[1].height = 150
    ws.freeze_panes = "F2"
    print(sname, "columnas", len(heads), "filas con datos", len({r['fila'] for r in rows}))

# Detalle
d = out.create_sheet("Detalle")
cols = ["Comunidad", "Año", "Estado", "Cuerpo", "Código", "Especialidad", "Idioma", "Plazas", "Libre",
        "Reserva discapacidad", "Acceso", "Ya aprobadas", "OPE extraordinaria", "Nota", "Fuente"]
for i, h in enumerate(cols, 1):
    c = d.cell(1, i, h)
    c.font = Font(name="Calibri", size=9, bold=True, color="FFFFFFFF")
    c.fill = PatternFill("solid", fgColor="FF1F4E79")
    c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = border
num = lambda s: int(s) if s and s.strip().lstrip("-").isdigit() else None
for i, r in enumerate(data, 2):
    vals = [r["fila"], num(r["anio"]), r["estado"], r["cuerpo"].zfill(4), r["codigo"], r["especialidad"], r["idioma"],
            num(r["plazas"]), num(r.get("libre", "")), num(r.get("reserva", "")), num(r.get("acceso", "")),
            num(r.get("aprobadas", "")), num(r.get("extra", "")), r["nota"], r["fuente"]]
    for j, v in enumerate(vals, 1):
        c = d.cell(i, j, v); c.font = Font(name="Calibri", size=9); c.border = border
for j, w in enumerate([18, 6, 16, 7, 6, 44, 10, 7, 6, 9, 7, 9, 10, 40, 50], 1):
    d.column_dimensions[get_column_letter(j)].width = w
d.freeze_panes = "A2"; d.auto_filter.ref = f"A1:O{len(data) + 1}"

n = out.create_sheet("Notas", 0)
notas = open(sys.argv[5], encoding="utf-8").read().strip().split("\n\n") if len(sys.argv) > 5 else []
n.column_dimensions["A"].width = 120
for i, t in enumerate(notas, 1):
    c = n.cell(i, 1, t); c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(name="Calibri", size=10)
out.active = 1
out.save(OUT)
